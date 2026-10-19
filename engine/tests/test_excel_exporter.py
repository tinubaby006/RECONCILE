import os
import tempfile
import openpyxl
from reconcile.export.excel_exporter import ExcelExporter, sanitize_cell_value
from reconcile.models import (
    ReconciliationRecord,
    ReconciliationSummary,
    RecordStatus,
    FieldDifference,
)


def test_sanitize_cell_value():
    # Formula injection attempts should be prefixed with apostrophe
    assert sanitize_cell_value("=1+1") == "'=1+1"
    assert sanitize_cell_value("+cmd|' /C calc'!A0") == "'+cmd|' /C calc'!A0"
    assert sanitize_cell_value("-5+10") == "'-5+10"
    assert sanitize_cell_value("@SUM(A1:A5)") == "'@SUM(A1:A5)"

    # Genuine numbers must NOT be sanitized or converted to strings
    assert sanitize_cell_value(100) == 100
    assert sanitize_cell_value(-50.25) == -50.25
    assert sanitize_cell_value(0.0) == 0.0

    # Normal strings
    assert sanitize_cell_value("Acme Corp") == "Acme Corp"
    assert sanitize_cell_value("INV-2024-001") == "INV-2024-001"
    assert sanitize_cell_value(None) == ""


def test_excel_export_structure_and_formatting():
    with tempfile.TemporaryDirectory() as tmp_dir:
        out_path = os.path.join(tmp_dir, "test_report.xlsx")

        summary = {
            "total_a": 3,
            "total_b": 3,
            "total_records": 4,
            "matched_count": 1,
            "discrepancies_count": 1,
            "missing_in_b_count": 1,
            "missing_in_a_count": 1,
            "duplicates_count": 0,
            "needs_review_count": 0,
            "total_variance": -100.0,
        }

        records = [
            {
                "id": "rec_a_2",
                "status": "matched",
                "status_label": "Matched Record",
                "entity": "Acme Corp",
                "reference": "INV-100",
                "source_a": "ledger.csv Row #2",
                "source_b": "bank.csv Row #2",
                "amount_a": 1000.0,
                "amount_b": 1000.0,
                "amount_diff": 0.0,
                "date_a": "2024-01-10",
                "date_b": "2024-01-10",
                "differences": [],
                "user_override": None,
            },
            {
                "id": "rec_a_3",
                "status": "discrepancy",
                "status_label": "Discrepancy",
                "entity": "Beta LLC",
                "reference": "=MALICIOUS_FORMULA",  # Injection test
                "source_a": "ledger.csv Row #3",
                "source_b": "bank.csv Row #3",
                "amount_a": 500.0,
                "amount_b": 400.0,
                "amount_diff": -100.0,
                "date_a": "2024-01-11",
                "date_b": "2024-01-12",
                "differences": [{"description": "Amount diff: -100.00"}],
                "user_override": None,
            },
            {
                "id": "rec_a_4",
                "status": "missing_in_b",
                "status_label": "Missing in B",
                "entity": "Gamma Inc",
                "reference": "INV-102",
                "source_a": "ledger.csv Row #4",
                "source_b": None,
                "amount_a": 250.0,
                "amount_b": None,
                "amount_diff": -250.0,
                "date_a": "2024-01-12",
                "date_b": None,
                "differences": [],
                "user_override": None,
            },
            {
                "id": "rec_b_4",
                "status": "missing_in_a",
                "status_label": "Missing in A",
                "entity": "Delta Ltd",
                "reference": "INV-999",
                "source_a": None,
                "source_b": "bank.csv Row #4",
                "amount_a": None,
                "amount_b": 300.0,
                "amount_diff": 300.0,
                "date_a": None,
                "date_b": "2024-01-15",
                "differences": [],
                "user_override": None,
            },
        ]

        saved_file = ExcelExporter.export(
            summary=summary,
            records=records,
            output_path=out_path,
            meta_a={"file_name": "ledger.csv", "total_rows": 3},
            meta_b={"file_name": "bank.csv", "total_rows": 3},
            scope="all",
        )

        assert os.path.exists(saved_file)
        assert os.path.getsize(saved_file) > 0

        # Open and verify workbook contents
        wb = openpyxl.load_workbook(saved_file, data_only=False)
        sheet_names = wb.sheetnames
        assert sheet_names == ["Summary", "Exceptions", "All Results", "By Entity"]

        # 1. Summary Sheet Checks
        ws_sum = wb["Summary"]
        assert ws_sum["A1"].value == "RECONCILIATION SUMMARY REPORT"
        assert ws_sum["A5"].value == "Dataset A (Primary)"
        assert ws_sum["B5"].value == "ledger.csv"
        assert ws_sum["A9"].value == "Status Metric"

        # 2. Exceptions Sheet Checks (Should only have the 3 exceptions)
        ws_exc = wb["Exceptions"]
        assert ws_exc.freeze_panes == "A2"
        assert ws_exc.auto_filter.ref is not None
        # Header + 3 exceptions = 4 rows
        assert ws_exc.max_row == 4
        # Verify formula injection was neutralized
        ref_cell = ws_exc.cell(row=2, column=3).value  # Second row, reference col
        assert ref_cell.startswith("'=") or ref_cell.startswith("INV-")

        # 3. All Results Sheet Checks (All 4 records)
        ws_all = wb["All Results"]
        assert ws_all.max_row == 5  # Header + 4 records

        # 4. By Entity Sheet Checks
        ws_ent = wb["By Entity"]
        assert ws_ent["A1"].value == "Entity / Company"
        # Check grand total row exists
        assert ws_ent.cell(row=ws_ent.max_row, column=1).value == "Grand Total"


def test_excel_export_exceptions_only_scope():
    with tempfile.TemporaryDirectory() as tmp_dir:
        out_path = os.path.join(tmp_dir, "test_exceptions_only.xlsx")

        summary = {"matched_count": 2, "discrepancies_count": 1, "total_records": 3}
        records = [
            {"id": "rec_1", "status": "matched", "reference": "M1", "entity": "A", "differences": []},
            {"id": "rec_2", "status": "matched", "reference": "M2", "entity": "B", "differences": []},
            {"id": "rec_3", "status": "discrepancy", "reference": "D1", "entity": "C", "differences": []},
        ]

        saved_file = ExcelExporter.export(
            summary=summary,
            records=records,
            output_path=out_path,
            scope="exceptions",
        )

        wb = openpyxl.load_workbook(saved_file)
        ws_all = wb["All Results"]
        # In exceptions-only scope, All Results only includes the 1 discrepancy
        # row 1 = header, row 2 = D1, plus optional scope note row
        assert ws_all.cell(row=2, column=3).value == "D1"


def test_excel_export_empty_records():
    with tempfile.TemporaryDirectory() as tmp_dir:
        out_path = os.path.join(tmp_dir, "empty_report.xlsx")
        summary = {"total_records": 0, "matched_count": 0}
        saved_file = ExcelExporter.export(
            summary=summary,
            records=[],
            output_path=out_path,
        )
        assert os.path.exists(saved_file)
        wb = openpyxl.load_workbook(saved_file)
        assert len(wb.sheetnames) == 4
