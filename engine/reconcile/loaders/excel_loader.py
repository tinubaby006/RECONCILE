from __future__ import annotations

import os
from typing import Optional, Tuple
import openpyxl
import pandas as pd

from reconcile.loaders.base import BaseLoader
from reconcile.models import DatasetMeta


class ExcelLoader(BaseLoader):
    """Loader for Excel workbooks (.xlsx, .xlsm)."""

    def inspect(self, file_path: str, **kwargs) -> DatasetMeta:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        wb = openpyxl.load_workbook(file_path, read_only=True, keep_links=False)
        sheet_names = wb.sheetnames
        wb.close()

        sheet_to_inspect = kwargs.get("sheet_name") or (sheet_names[0] if sheet_names else None)
        meta, _ = self.load(file_path, sheet_name=sheet_to_inspect, nrows=50)
        meta.sheets = sheet_names
        meta.active_sheet = sheet_to_inspect
        return meta

    def load(
        self,
        file_path: str,
        sheet_name: Optional[str] = None,
        nrows: Optional[int] = None,
        **kwargs,
    ) -> Tuple[DatasetMeta, pd.DataFrame]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        wb = openpyxl.load_workbook(file_path, read_only=True, keep_links=False)
        all_sheets = wb.sheetnames
        wb.close()

        target_sheet = sheet_name or (all_sheets[0] if all_sheets else 0)

        df = pd.read_excel(
            file_path,
            sheet_name=target_sheet,
            nrows=nrows,
            dtype=str,
            keep_default_na=False,
            engine="openpyxl",
        )

        resolved_sheet_name = (
            target_sheet if isinstance(target_sheet, str) else all_sheets[target_sheet]
        )

        file_name = os.path.basename(file_path)
        # Header is row 1, first data row is Excel row 2
        df["_source_file"] = file_name
        df["_source_row"] = (df.index + 2).astype(int)
        df["_source_sheet"] = resolved_sheet_name

        meta = self.build_metadata(
            file_path=file_path,
            file_type="xlsx",
            df=df,
            sheets=all_sheets,
            active_sheet=resolved_sheet_name,
        )
        return meta, df
