from __future__ import annotations

import datetime
import os
import re
from typing import Any, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def sanitize_cell_value(val: Any) -> Any:
    """Neutralizes formula injection in untrusted string values beginning with =, +, -, or @.

    Genuine numeric values (integers, floats) are preserved as real numbers.
    """
    if val is None:
        return ""

    # Preserve genuine numbers and dates
    if isinstance(val, (int, float, datetime.date, datetime.datetime)):
        return val

    # Convert to string and sanitize formulas
    s = str(val)
    if s and s[0] in ("=", "+", "-", "@"):
        # Prefix with a single apostrophe so Excel renders it as plain text
        return "'" + s
    return s


class ExcelExporter:
    """Generates production-grade multi-sheet Excel workbooks for reconciliation results using openpyxl."""

    HEADER_FILL = PatternFill(start_color="1A146B", end_color="1A146B", fill_type="solid")
    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    SUBHEADER_FILL = PatternFill(start_color="E9EDFF", end_color="E9EDFF", fill_type="solid")
    SUBHEADER_FONT = Font(name="Calibri", size=11, bold=True, color="141B2B")
    TOTAL_FONT = Font(name="Calibri", size=11, bold=True, color="141B2B")
    BORDER_THIN = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )
    BORDER_TOP_THIN_BOTTOM_DOUBLE = Border(
        top=Side(style="thin", color="A0A0A0"),
        bottom=Side(style="double", color="141B2B"),
    )

    STATUS_COLORS = {
        "matched": PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid"),
        "discrepancy": PatternFill(start_color="FFF1F2", end_color="FFF1F2", fill_type="solid"),
        "needs_review": PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid"),
        "missing_in_b": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
        "missing_in_a": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
        "duplicate": PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid"),
    }

    @classmethod
    def export(
        cls,
        summary: dict[str, Any],
        records: list[dict[str, Any]],
        output_path: str,
        meta_a: Optional[dict[str, Any]] = None,
        meta_b: Optional[dict[str, Any]] = None,
        scope: str = "exceptions",
        overwrite: bool = True,
    ) -> str:
        """Builds and saves the reconciliation Excel workbook.

        Returns the verified absolute path of the generated file.
        """
        # Validate output directory
        target_dir = os.path.dirname(os.path.abspath(output_path))
        if not os.path.exists(target_dir):
            os.makedirs(target_dir, exist_ok=True)

        # Normalize filename
        base_name = os.path.basename(output_path)
        if not base_name.lower().endswith(".xlsx"):
            base_name += ".xlsx"
        
        # Clean illegal characters in Windows filename
        base_name = re.sub(r'[\\/*?:"<>|]', "_", base_name)
        final_path = os.path.join(target_dir, base_name)

        if os.path.exists(final_path) and not overwrite:
            # Generate unique filename if not overwriting
            stem, ext = os.path.splitext(base_name)
            counter = 1
            while os.path.exists(os.path.join(target_dir, f"{stem}_{counter}{ext}")):
                counter += 1
            final_path = os.path.join(target_dir, f"{stem}_{counter}{ext}")

        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # 1. Summary Sheet
        cls._build_summary_sheet(wb, summary, meta_a, meta_b, scope, total_recs=len(records))

        # 2. Exceptions Sheet
        cls._build_records_sheet(wb, "Exceptions", records, filter_exceptions=True)

        # 3. All Results Sheet
        cls._build_records_sheet(wb, "All Results", records, filter_exceptions=(scope == "exceptions"), scope_note=(scope == "exceptions"))

        # 4. By Entity Sheet
        cls._build_by_entity_sheet(wb, records)

        # Save workbook
        wb.save(final_path)

        # Verify file exists and is non-empty
        if not os.path.exists(final_path) or os.path.getsize(final_path) == 0:
            raise IOError(f"Failed to verify written Excel file at: {final_path}")

        return os.path.abspath(final_path)

    @classmethod
    def _build_summary_sheet(
        cls,
        wb: openpyxl.Workbook,
        summary: dict[str, Any],
        meta_a: Optional[dict[str, Any]],
        meta_b: Optional[dict[str, Any]],
        scope: str,
        total_recs: int,
    ) -> None:
        ws = wb.create_sheet(title="Summary")
        ws.views.sheetView[0].showGridLines = True

        # Report Header
        ws["A1"] = "RECONCILIATION SUMMARY REPORT"
        ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="1A146B")

        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws["A2"] = f"Generated locally on: {now_str} | Scope: {scope.capitalize()}"
        ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="555555")

        # Section 1: Datasets Information
        ws["A4"] = "Input Datasets"
        ws["A4"].font = cls.SUBHEADER_FONT

        ws["A5"] = "Dataset A (Primary)"
        ws["B5"] = sanitize_cell_value(meta_a.get("file_name") if meta_a else "Dataset A")
        ws["C5"] = f"{meta_a.get('total_rows', 0)} rows" if meta_a else ""

        ws["A6"] = "Dataset B (Comparison)"
        ws["B6"] = sanitize_cell_value(meta_b.get("file_name") if meta_b else "Dataset B")
        ws["C6"] = f"{meta_b.get('total_rows', 0)} rows" if meta_b else ""

        for row in range(5, 7):
            ws[f"A{row}"].font = Font(bold=True)
            for col in ("A", "B", "C"):
                ws[f"{col}{row}"].border = cls.BORDER_THIN

        # Section 2: Metrics Breakdown
        ws["A8"] = "Reconciliation Results"
        ws["A8"].font = cls.SUBHEADER_FONT

        headers = ["Status Metric", "Record Count", "Percentage of Total"]
        for c_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=9, column=c_idx, value=h)
            cell.font = cls.HEADER_FONT
            cell.fill = cls.HEADER_FILL
            cell.alignment = Alignment(horizontal="center" if c_idx > 1 else "left")

        tot_records = summary.get("total_records", total_recs)
        if tot_records == 0:
            tot_records = max(1, total_recs)

        metrics = [
            ("Confirmed Matched", summary.get("matched_count", 0)),
            ("Selected-Field Discrepancies", summary.get("discrepancies_count", 0)),
            ("Requires Review (Ambiguous / Fuzzy)", summary.get("needs_review_count", 0)),
            ("Missing in Dataset B", summary.get("missing_in_b_count", 0)),
            ("Missing in Dataset A", summary.get("missing_in_a_count", 0)),
            ("Duplicate Records", summary.get("duplicates_count", 0)),
        ]

        curr_row = 10
        for label, cnt in metrics:
            ws.cell(row=curr_row, column=1, value=label).border = cls.BORDER_THIN
            c_cnt = ws.cell(row=curr_row, column=2, value=cnt)
            c_cnt.border = cls.BORDER_THIN
            c_cnt.number_format = "#,##0"
            c_cnt.alignment = Alignment(horizontal="right")

            pct = (cnt / tot_records) if tot_records > 0 else 0.0
            c_pct = ws.cell(row=curr_row, column=3, value=pct)
            c_pct.border = cls.BORDER_THIN
            c_pct.number_format = "0.0%"
            c_pct.alignment = Alignment(horizontal="right")
            curr_row += 1

        # Totals Row
        tot_row = curr_row
        ws.cell(row=tot_row, column=1, value="Total Records Evaluated").font = cls.TOTAL_FONT
        ws.cell(row=tot_row, column=1).border = cls.BORDER_TOP_THIN_BOTTOM_DOUBLE

        c_tot = ws.cell(row=tot_row, column=2, value=f"=SUM(B10:B{tot_row-1})")
        c_tot.font = cls.TOTAL_FONT
        c_tot.number_format = "#,##0"
        c_tot.alignment = Alignment(horizontal="right")
        c_tot.border = cls.BORDER_TOP_THIN_BOTTOM_DOUBLE

        c_tot_pct = ws.cell(row=tot_row, column=3, value=f"=SUM(C10:C{tot_row-1})")
        c_tot_pct.font = cls.TOTAL_FONT
        c_tot_pct.number_format = "0.0%"
        c_tot_pct.alignment = Alignment(horizontal="right")
        c_tot_pct.border = cls.BORDER_TOP_THIN_BOTTOM_DOUBLE

        # Net Variance Summary
        curr_row += 2
        ws.cell(row=curr_row, column=1, value="Net Amount Variance").font = cls.TOTAL_FONT
        c_var = ws.cell(row=curr_row, column=2, value=float(summary.get("total_variance", 0.0)))
        c_var.font = cls.TOTAL_FONT
        c_var.number_format = "$#,##0.00;($#,##0.00);\"-\""
        c_var.alignment = Alignment(horizontal="right")

        # Auto-adjust column widths
        ws.column_dimensions["A"].width = 38
        ws.column_dimensions["B"].width = 24
        ws.column_dimensions["C"].width = 24

    @classmethod
    def _build_records_sheet(
        cls,
        wb: openpyxl.Workbook,
        title: str,
        records: list[dict[str, Any]],
        filter_exceptions: bool = False,
        scope_note: bool = False,
    ) -> None:
        ws = wb.create_sheet(title=title)
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A2"

        headers = [
            "Record ID",
            "Status",
            "Reference Key",
            "Entity / Company",
            "Source A",
            "Source B",
            "Amount A",
            "Amount B",
            "Variance (Delta)",
            "Date A",
            "Date B",
            "Differences & Notes",
            "Override Status",
        ]

        # Write Header
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = cls.HEADER_FONT
            cell.fill = cls.HEADER_FILL
            cell.alignment = Alignment(horizontal="center" if "Amount" in h or "Variance" in h else "left")

        # Filter records if requested
        if filter_exceptions:
            row_items = [r for r in records if r.get("status") != "matched"]
        else:
            row_items = records

        row_num = 2
        for r in row_items:
            status = r.get("status", "")
            diff_notes = "; ".join(
                d.get("description", "") for d in r.get("differences", [])
            )

            row_data = [
                sanitize_cell_value(r.get("id")),
                sanitize_cell_value(r.get("status_label") or status),
                sanitize_cell_value(r.get("reference")),
                sanitize_cell_value(r.get("entity")),
                sanitize_cell_value(r.get("source_a")),
                sanitize_cell_value(r.get("source_b")),
                r.get("amount_a"),  # Real numeric float/int
                r.get("amount_b"),  # Real numeric float/int
                r.get("amount_diff"),  # Real numeric float/int
                sanitize_cell_value(r.get("date_a")),
                sanitize_cell_value(r.get("date_b")),
                sanitize_cell_value(diff_notes),
                sanitize_cell_value(r.get("user_override") or ""),
            ]

            status_fill = cls.STATUS_COLORS.get(status)

            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=row_num, column=col_idx, value=val)
                cell.border = cls.BORDER_THIN

                # Status cell highlighting
                if col_idx == 2 and status_fill:
                    cell.fill = status_fill

                # Formatting for numeric amounts
                if col_idx in (7, 8, 9):
                    if isinstance(val, (int, float)):
                        cell.number_format = "#,##0.00;(#,##0.00);\"-\""
                        cell.alignment = Alignment(horizontal="right")
                    else:
                        cell.alignment = Alignment(horizontal="center")

            row_num += 1

        # Enable auto-filters
        last_col_letter = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col_letter}{max(row_num-1, 1)}"

        # Add scope note at the bottom if filtered
        if scope_note and filter_exceptions:
            note_row = row_num + 1
            ws.cell(row=note_row, column=1, value="Note: Sheet filtered to exception records only based on user export selection.").font = Font(italic=True, color="777777")

        # Auto-fit column widths
        cls._autofit_columns(ws, max_cols=len(headers))

    @classmethod
    def _build_by_entity_sheet(cls, wb: openpyxl.Workbook, records: list[dict[str, Any]]) -> None:
        ws = wb.create_sheet(title="By Entity")
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A2"

        headers = [
            "Entity / Company",
            "Total Records",
            "Matched",
            "Discrepancies",
            "Needs Review",
            "Missing in B",
            "Missing in A",
            "Total Amount A",
            "Total Amount B",
            "Net Variance",
        ]

        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = cls.HEADER_FONT
            cell.fill = cls.HEADER_FILL
            cell.alignment = Alignment(horizontal="right" if col_idx > 1 else "left")

        # Group records by entity
        grouped: dict[str, list[dict[str, Any]]] = {}
        for r in records:
            ent = r.get("entity") or "Unassigned"
            grouped.setdefault(ent, []).append(r)

        row_num = 2
        for ent_name, ent_records in sorted(grouped.items()):
            tot_cnt = len(ent_records)
            matched = sum(1 for r in ent_records if r.get("status") == "matched")
            discrepancies = sum(1 for r in ent_records if r.get("status") == "discrepancy")
            needs_review = sum(1 for r in ent_records if r.get("status") == "needs_review")
            missing_b = sum(1 for r in ent_records if r.get("status") == "missing_in_b")
            missing_a = sum(1 for r in ent_records if r.get("status") == "missing_in_a")

            sum_amt_a = sum(r.get("amount_a") or 0.0 for r in ent_records if r.get("amount_a") is not None)
            sum_amt_b = sum(r.get("amount_b") or 0.0 for r in ent_records if r.get("amount_b") is not None)
            net_diff = sum_amt_b - sum_amt_a

            row_vals = [
                sanitize_cell_value(ent_name),
                tot_cnt,
                matched,
                discrepancies,
                needs_review,
                missing_b,
                missing_a,
                round(sum_amt_a, 2),
                round(sum_amt_b, 2),
                round(net_diff, 2),
            ]

            for c_idx, val in enumerate(row_vals, start=1):
                cell = ws.cell(row=row_num, column=c_idx, value=val)
                cell.border = cls.BORDER_THIN
                if c_idx >= 8:
                    cell.number_format = "#,##0.00;(#,##0.00);\"-\""
                    cell.alignment = Alignment(horizontal="right")
                elif c_idx > 1:
                    cell.number_format = "#,##0"
                    cell.alignment = Alignment(horizontal="right")

            row_num += 1

        # Totals Row at the bottom
        if row_num > 2:
            ws.cell(row=row_num, column=1, value="Grand Total").font = cls.TOTAL_FONT
            ws.cell(row=row_num, column=1).border = cls.BORDER_TOP_THIN_BOTTOM_DOUBLE

            for col_idx in range(2, 8):
                c_letter = get_column_letter(col_idx)
                c_cell = ws.cell(row=row_num, column=col_idx, value=f"=SUM({c_letter}2:{c_letter}{row_num-1})")
                c_cell.font = cls.TOTAL_FONT
                c_cell.number_format = "#,##0"
                c_cell.alignment = Alignment(horizontal="right")
                c_cell.border = cls.BORDER_TOP_THIN_BOTTOM_DOUBLE

            for col_idx in (8, 9, 10):
                c_letter = get_column_letter(col_idx)
                c_cell = ws.cell(row=row_num, column=col_idx, value=f"=SUM({c_letter}2:{c_letter}{row_num-1})")
                c_cell.font = cls.TOTAL_FONT
                c_cell.number_format = "#,##0.00;(#,##0.00);\"-\""
                c_cell.alignment = Alignment(horizontal="right")
                c_cell.border = cls.BORDER_TOP_THIN_BOTTOM_DOUBLE

        last_col_letter = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col_letter}{max(row_num-1, 1)}"
        cls._autofit_columns(ws, max_cols=len(headers))

    @staticmethod
    def _autofit_columns(ws: openpyxl.worksheet.worksheet.Worksheet, max_cols: int) -> None:
        for col_idx in range(1, max_cols + 1):
            col_letter = get_column_letter(col_idx)
            max_len = 0
            for row in range(1, min(ws.max_row + 1, 500)):
                val = ws.cell(row=row, column=col_idx).value
                if val:
                    s_val = str(val)
                    if len(s_val) > max_len:
                        max_len = len(s_val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
